from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
F = "Arial"
BLUE = Font(name=F, color="0000FF", size=10)
BLK = Font(name=F, size=10)
GRN = Font(name=F, color="008000", size=10)
BOLD = Font(name=F, bold=True, size=10)
HDR = Font(name=F, bold=True, color="FFFFFF", size=10)
TITLE = Font(name=F, bold=True, size=14)
HFILL = PatternFill("solid", fgColor="1F3864")
YEL = PatternFill("solid", fgColor="FFFF00")
TOT = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BBBBBB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '$#,##0;($#,##0);-'
PCT = '0.0%;(0.0%);-'

def put(ws, ref, v, font=BLK, fmt=None, fill=None, bold=False):
    c = ws[ref]; c.value = v
    c.font = Font(name=font.name, size=font.size, color=font.color, bold=bold or font.bold)
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    c.border = BOX
    return c

def header(ws, row, labels, col0=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col0+i, value=l)
        c.font = HDR; c.fill = HFILL; c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 32

def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

# ---------------- Inputs ----------------
ws = wb.active; ws.title = "Inputs"
ws["A1"] = "Lab 3b - Inputs"; ws["A1"].font = TITLE
ws["A2"] = "Blue = given in the lab dataset. Yellow fill = my assumption (edit to run scenarios). Black = formula."
ws["A2"].font = Font(name=F, italic=True, size=9)
ws["A3"] = "System context (SecureMobile)"; ws["A3"].font = BOLD
for r, (l, v, f) in enumerate([("Active customers", 500000, '#,##0'), ("Average transaction ($)", 2500, USD),
                               ("Daily transactions", 50000, '#,##0')], 4):
    put(ws, f"A{r}", l); put(ws, f"B{r}", v, BLUE, f)

header(ws, 8, ["Finding", "Exploit probability (= ARO)", "Loss quantity", "Unit loss ($)", "SLE basis", "ID"])
rows = [("API Authentication Bypass", 0.15, 1, 5000000, "1 incident x max single-incident loss", "F1"),
        ("Database Injection", 0.25, "=B4", 250, "records at risk x cost per record", "F2"),
        ("Session Hijacking", 0.40, 5000, 1500, "accounts at risk x avg loss per account", "F3")]
for i, (n, p, q, u, b, k) in enumerate(rows):
    r = 9+i
    put(ws, f"A{r}", n, BLUE); put(ws, f"B{r}", p, BLUE, '0%')
    put(ws, f"C{r}", q, BLK if isinstance(q, str) else BLUE, '#,##0')
    put(ws, f"D{r}", u, BLUE, USD); put(ws, f"E{r}", b); put(ws, f"F{r}", k)

header(ws, 13, ["Control", "Initial cost ($)", "Annual maintenance ($)", "Effectiveness", "Mitigates (ID)"])
ctrl = [("Advanced API Security Gateway", 350000, 50000, 0.90, "F1"),
        ("Web Application Firewall (WAF)", 150000, 25000, 0.75, "F2"),
        ("Multi-Factor Authentication Enhancement", 200000, 30000, 0.95, "F3")]
for i, (n, a, m, e, k) in enumerate(ctrl):
    r = 14+i
    put(ws, f"A{r}", n, BLUE); put(ws, f"B{r}", a, BLUE, USD); put(ws, f"C{r}", m, BLUE, USD)
    put(ws, f"D{r}", e, BLUE, '0%'); put(ws, f"E{r}", k, BLUE, fill=YEL)

ws["A18"] = "Assumptions and thresholds (all mine - not in the lab brief)"; ws["A18"].font = BOLD
assump = [("Evaluation horizon (years)", 3, '0'),
          ("Priority 'Critical' if ALE >=", 10000000, USD), ("Priority 'High' if ALE >=", 1000000, USD),
          ("Priority 'Medium' if ALE >=", 250000, USD),
          ("Probability band 'High' if >=", 0.35, '0%'), ("Probability band 'Medium' if >=", 0.20, '0%'),
          ("Impact band 'High' if SLE >=", 10000000, USD), ("Impact band 'Medium' if SLE >=", 1000000, USD),
          ("Bonus: budget cap on Year-1 outlay ($)", 450000, USD)]
for i, (l, v, f) in enumerate(assump):
    r = 19+i
    put(ws, f"A{r}", l); put(ws, f"B{r}", v, BLUE, f, YEL)

ws["A29"] = "Heat-map rating matrix (impact down, probability across)"; ws["A29"].font = BOLD
header(ws, 30, ["Impact \\ Probability", "Low", "Medium", "High"])
for i, (imp, vals) in enumerate([("High", ["Medium", "High", "High"]), ("Medium", ["Medium", "Medium", "High"]),
                                 ("Low", ["Low", "Low", "Medium"])]):
    put(ws, f"A{31+i}", imp, BOLD)
    for j, v in enumerate(vals):
        put(ws, f"{'BCD'[j]}{31+i}", v, BLUE, fill=YEL)

notes = ["Notes:",
 "1. ARO = exploit probability per year (the lab gives probability only; I treat it as one expected exploitation attempt-year).",
 "2. F1 SLE uses the 'Maximum Single Incident Loss' given. The 48-hour detection time is context only; it is not used in ALE.",
 "3. F2 SLE assumes all 500,000 records are exposed (worst case) at $250/record.",
 "4. Each control is mapped to the single finding it addresses: API Gateway->F1, WAF->F2 (injection), MFA->F3 (account takeover).",
 "5. Effectiveness is applied as a straight % reduction of that finding's ALE."]
for i, n in enumerate(notes):
    ws[f"A{36+i}"] = n; ws[f"A{36+i}"].font = Font(name=F, size=9, bold=(i == 0))
widths(ws, [44, 22, 18, 16, 40, 8])

# ---------------- Risk ----------------
rk = wb.create_sheet("Risk")
rk["A1"] = "Phase 1 - Risk Exposure (SLE x ARO = ALE)"; rk["A1"].font = TITLE
header(rk, 4, ["Vulnerability", "SLE ($)", "ARO", "ALE ($)", "% of total ALE", "Risk priority", "Probability band", "Impact band", "Heat-map rating", "Rank"])
for i in range(3):
    r = 5+i; s = 9+i
    put(rk, f"A{r}", f"=Inputs!A{s}", GRN)
    put(rk, f"B{r}", f"=Inputs!C{s}*Inputs!D{s}", GRN, USD)
    put(rk, f"C{r}", f"=Inputs!B{s}", GRN, '0%')
    put(rk, f"D{r}", f"=B{r}*C{r}", BLK, USD)
    put(rk, f"E{r}", f"=D{r}/$D$8", BLK, PCT)
    put(rk, f"F{r}", f'=IF(D{r}>=Inputs!$B$20,"Critical",IF(D{r}>=Inputs!$B$21,"High",IF(D{r}>=Inputs!$B$22,"Medium","Low")))')
    put(rk, f"G{r}", f'=IF(C{r}>=Inputs!$B$23,"High",IF(C{r}>=Inputs!$B$24,"Medium","Low"))')
    put(rk, f"H{r}", f'=IF(B{r}>=Inputs!$B$25,"High",IF(B{r}>=Inputs!$B$26,"Medium","Low"))')
    put(rk, f"I{r}", f"=INDEX(Inputs!$B$31:$D$33,MATCH(H{r},Inputs!$A$31:$A$33,0),MATCH(G{r},Inputs!$B$30:$D$30,0))")
    put(rk, f"J{r}", f"=RANK(D{r},$D$5:$D$7)", BLK, '0')
put(rk, "A8", "Total", BOLD, fill=TOT)
for c in "BCFGHIJ": put(rk, f"{c}8", None, fill=TOT)
put(rk, "D8", "=SUM(D5:D7)", BOLD, USD, TOT); put(rk, "E8", "=SUM(E5:E7)", BOLD, PCT, TOT)
widths(rk, [34, 18, 8, 18, 14, 14, 16, 14, 16, 8])

# ---------------- Controls ----------------
ct = wb.create_sheet("Controls")
ct["A1"] = "Phase 2 - Cost-Benefit Analysis"; ct["A1"].font = TITLE
ct["A2"] = "Year-1 ROI = (annual risk reduction - (initial + annual maintenance)) / (initial + annual maintenance). Payback = initial / (annual risk reduction - annual maintenance)."
ct["A2"].font = Font(name=F, italic=True, size=9)
header(ct, 4, ["Control", "Initial investment ($)", "Annual maintenance ($)", "Mitigates", "Baseline ALE ($)", "Effectiveness",
               "Annual risk reduction ($)", "Residual ALE ($)", "Year-1 cost ($)", "Year-1 ROI",
               "Horizon cost ($)", "Horizon benefit ($)", "Horizon ROI", "Payback (years)", "Payback (days)"])
for i in range(3):
    r = 5+i; s = 14+i
    put(ct, f"A{r}", f"=Inputs!A{s}", GRN)
    put(ct, f"B{r}", f"=Inputs!B{s}", GRN, USD); put(ct, f"C{r}", f"=Inputs!C{s}", GRN, USD)
    put(ct, f"D{r}", f"=Inputs!E{s}", GRN)
    put(ct, f"E{r}", f"=INDEX(Risk!$D$5:$D$7,MATCH(D{r},Inputs!$F$9:$F$11,0))", GRN, USD)
    put(ct, f"F{r}", f"=Inputs!D{s}", GRN, '0%')
    put(ct, f"G{r}", f"=E{r}*F{r}", BLK, USD); put(ct, f"H{r}", f"=E{r}-G{r}", BLK, USD)
    put(ct, f"I{r}", f"=B{r}+C{r}", BLK, USD); put(ct, f"J{r}", f"=(G{r}-I{r})/I{r}", BLK, '#,##0.0%')
    put(ct, f"K{r}", f"=B{r}+C{r}*Inputs!$B$19", GRN, USD); put(ct, f"L{r}", f"=G{r}*Inputs!$B$19", GRN, USD)
    put(ct, f"M{r}", f"=(L{r}-K{r})/K{r}", BLK, '#,##0.0%')
    put(ct, f"N{r}", f"=B{r}/(G{r}-C{r})", BLK, '0.000'); put(ct, f"O{r}", f"=N{r}*365", BLK, '0.0')
put(ct, "A8", "Portfolio (all three)", BOLD, fill=TOT)
for c in "DF": put(ct, f"{c}8", None, fill=TOT)
for c in "BCEGHIKL": put(ct, f"{c}8", f"=SUM({c}5:{c}7)", BOLD, USD, TOT)
put(ct, "J8", "=(G8-I8)/I8", BOLD, '#,##0.0%', TOT); put(ct, "M8", "=(L8-K8)/K8", BOLD, '#,##0.0%', TOT)
put(ct, "N8", "=B8/(G8-C8)", BOLD, '0.000', TOT); put(ct, "O8", "=N8*365", BOLD, '0.0', TOT)
widths(ct, [38, 16, 16, 10, 16, 12, 18, 16, 14, 12, 14, 16, 12, 12, 12])

# ---------------- Summary ----------------
sm = wb.create_sheet("Summary")
sm["A1"] = "Executive Summary - Before vs After Controls"; sm["A1"].font = TITLE
header(sm, 4, ["Vulnerability", "ALE before ($)", "Control", "Residual ALE ($)", "Risk reduction ($)", "Reduction %"])
for i in range(3):
    r = 5+i
    put(sm, f"A{r}", f"=Risk!A{r}", GRN); put(sm, f"B{r}", f"=Risk!D{r}", GRN, USD)
    put(sm, f"C{r}", f"=Controls!A{r}", GRN); put(sm, f"D{r}", f"=Controls!H{r}", GRN, USD)
    put(sm, f"E{r}", f"=B{r}-D{r}", BLK, USD); put(sm, f"F{r}", f"=E{r}/B{r}", BLK, PCT)
put(sm, "A8", "Total", BOLD, fill=TOT); put(sm, "C8", None, fill=TOT)
for c in "BDE": put(sm, f"{c}8", f"=SUM({c}5:{c}7)", BOLD, USD, TOT)
put(sm, "F8", "=E8/B8", BOLD, PCT, TOT)
kp = [("Total initial investment", "=Controls!B8", USD), ("Annual maintenance", "=Controls!C8", USD),
      ("Year-1 total cost", "=Controls!I8", USD), ("Annual risk reduction", "=Controls!G8", USD),
      ("Portfolio Year-1 ROI", "=Controls!J8", '#,##0.0%'), ("Portfolio payback (days)", "=Controls!O8", '0.0'),
      ("Residual ALE after all controls", "=D8", USD)]
header(sm, 11, ["KPI", "Value"])
for i, (l, f, fm) in enumerate(kp):
    put(sm, f"A{12+i}", l); put(sm, f"B{12+i}", f, GRN, fm)
widths(sm, [38, 18, 40, 18, 18, 14])

# ---------------- Projection ----------------
pj = wb.create_sheet("Projection")
pj["A1"] = "Risk reduction over time"; pj["A1"].font = TITLE
pj["A3"] = "Cumulative net benefit by year ($) = -initial + (annual risk reduction - annual maintenance) x year"; pj["A3"].font = BOLD
header(pj, 4, ["Control"] + [f"Year {y}" for y in range(6)])
put(pj, "A5", "Year index", BOLD)
for y in range(6): put(pj, f"{get_column_letter(2+y)}5", y, BLUE, '0')
for i in range(3):
    r = 6+i; put(pj, f"A{r}", f"=Controls!A{5+i}", GRN)
    for y in range(6):
        c = get_column_letter(2+y)
        put(pj, f"{c}{r}", f"=-Controls!$B${5+i}+(Controls!$G${5+i}-Controls!$C${5+i})*{c}$5", BLK, USD)
put(pj, "A9", "Portfolio", BOLD, fill=TOT)
for y in range(6):
    c = get_column_letter(2+y); put(pj, f"{c}9", f"=SUM({c}6:{c}8)", BOLD, USD, TOT)

pj["A12"] = "Phased roll-out: go-live month per control (assumption)"; pj["A12"].font = BOLD
header(pj, 13, ["Control", "Go-live month"])
for i, m in enumerate([6, 1, 3]):
    put(pj, f"A{14+i}", f"=Controls!A{5+i}", GRN); put(pj, f"B{14+i}", m, BLUE, '0', YEL)
header(pj, 18, ["Month", "Annualised residual exposure ($)"])
for m in range(13):
    r = 19+m; put(pj, f"A{r}", m, BLUE, '0')
    put(pj, f"B{r}", f"=Risk!$D$8-SUMPRODUCT((A{r}>=$B$14:$B$16)*Controls!$G$5:$G$7)", BLK, USD)
widths(pj, [44, 22, 16, 16, 16, 16, 16])

# ---------------- Bonus ----------------
bn = wb.create_sheet("Bonus_Budget")
bn["A1"] = "Bonus - optimal control selection under a budget"; bn["A1"].font = TITLE
put(bn, "A3", "Budget cap on Year-1 outlay ($)", BOLD); put(bn, "B3", "=Inputs!B27", GRN, USD)
header(bn, 5, ["Combination", "API GW", "WAF", "MFA", "Initial ($)", "Year-1 outlay ($)", "Annual risk reduction ($)", "Year-1 net benefit ($)", "Within budget?", "Eligible net benefit ($)"])
combos = [("None", 0, 0, 0), ("API GW", 1, 0, 0), ("WAF", 0, 1, 0), ("MFA", 0, 0, 1), ("API GW + WAF", 1, 1, 0),
          ("API GW + MFA", 1, 0, 1), ("WAF + MFA", 0, 1, 1), ("All three", 1, 1, 1)]
for i, (n, a, w, m) in enumerate(combos):
    r = 6+i
    put(bn, f"A{r}", n); put(bn, f"B{r}", a, BLUE, '0'); put(bn, f"C{r}", w, BLUE, '0'); put(bn, f"D{r}", m, BLUE, '0')
    put(bn, f"E{r}", f"=B{r}*Controls!$B$5+C{r}*Controls!$B$6+D{r}*Controls!$B$7", GRN, USD)
    put(bn, f"F{r}", f"=B{r}*Controls!$I$5+C{r}*Controls!$I$6+D{r}*Controls!$I$7", GRN, USD)
    put(bn, f"G{r}", f"=B{r}*Controls!$G$5+C{r}*Controls!$G$6+D{r}*Controls!$G$7", GRN, USD)
    put(bn, f"H{r}", f"=G{r}-F{r}", BLK, USD)
    put(bn, f"I{r}", f'=IF(F{r}<=$B$3,"Yes","No")')
    put(bn, f"J{r}", f'=IF(I{r}="Yes",H{r},"")', BLK, USD)
put(bn, "A16", "Best affordable combination", BOLD, fill=TOT)
put(bn, "B16", "=INDEX(A6:A13,MATCH(MAX(J6:J13),J6:J13,0))", BOLD, fill=TOT)
put(bn, "A17", "Its Year-1 net benefit ($)", BOLD, fill=TOT); put(bn, "B17", "=MAX(J6:J13)", BOLD, USD, TOT)
widths(bn, [30, 12, 10, 10, 16, 18, 22, 22, 14, 22])

wb.save("/mnt/user-data/outputs/risk-lab/Lab3b_Risk_Quantification.xlsx")
